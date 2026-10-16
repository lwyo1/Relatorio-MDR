"""Teste de ouro: o relatório gerado para o MDR-001 deve reproduzir o modelo do Gripen.

Compara, seção a seção, as tabelas do .docx gerado com as do modelo preenchido
(modelos/modelo_gripen_referencia.docx). As únicas diferenças aceitas são as
decididas no planejamento e listadas em DIFERENCAS_ESPERADAS.
"""

import re
import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import docx
from docx.table import Table
from docx.text.paragraph import Paragraph

from mdr import gerador, montagem, textos
from tests.dados_gripen import base_gripen

RAIZ = Path(__file__).resolve().parent.parent
REFERENCIA = RAIZ / "modelos" / "modelo_gripen_referencia.docx"
TEMPLATE = RAIZ / "modelos" / "relatorio_mdr.docx"

DIFERENCAS_ESPERADAS = """
- Data de geração: automática no sistema; no modelo é texto explicativo.
- 'Parceiro principal' passa a 'Parceiro(s)' (D12).
- Seção 3: tabela vertical de 2 colunas com os campos completos (D5).
- Seção 4: Leitura gerencial vem do template por grau (item 4.5).
- Seção 6: sem 'Variedade observável' no relatório individual (D11).
- Seção 8: rótulo genérico 'Trajetória da função no período/fase analisada'.
- Seção 12: 'Condição de preenchimento' usa o texto fixo genérico (item 12.4).
"""


def _norm(texto: str) -> str:
    return " ".join(texto.split()).rstrip(".").strip()


def _tabelas_por_secao(caminho: Path) -> dict[str, list[Table]]:
    """{'4': [tabela, caixa, ...]}: tabelas que vêm após cada título numerado."""
    d = docx.Document(str(caminho))
    secoes: dict[str, list[Table]] = {}
    atual = "0"
    for el in d.element.body.iterchildren():
        if el.tag.endswith("}p"):
            p = Paragraph(el, d)
            m = re.match(r"(\d+)\.\s", p.text)
            if p.style.name.startswith("Heading") and m:
                atual = m.group(1)
        elif el.tag.endswith("}tbl"):
            secoes.setdefault(atual, []).append(Table(el, d))
    return secoes


def _linhas(tabela: Table) -> list[tuple[str, ...]]:
    return [tuple(_norm(c.text) for c in row.cells) for row in tabela.rows]


def _caixa(tabela: Table) -> str:
    texto = tabela.cell(0, 0).text
    return _norm(texto.split("\n", 1)[1])


class TesteDeOuroGripen(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        base = base_gripen()
        conteudo = montagem.montar(base.registro("MDR-001"), base,
                                   agora=datetime(2026, 10, 15, 21, 0))
        cls.tmp = tempfile.TemporaryDirectory()
        cls.saida = gerador.gerar(TEMPLATE, conteudo, Path(cls.tmp.name) / "r.docx")
        cls.ger = _tabelas_por_secao(cls.saida)
        cls.ref = _tabelas_por_secao(REFERENCIA)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def comparar(self, secao, idx=0, ignorar_rotulos=(), so_valores=False):
        ger, ref = _linhas(self.ger[secao][idx]), _linhas(self.ref[secao][idx])
        ref = [r for r in ref if r[0] not in ignorar_rotulos]
        ger = [g for g in ger if g[0] not in ignorar_rotulos]
        if so_valores:
            ger, ref = [g[1:] for g in ger], [r[1:] for r in ref]
        self.assertEqual(ger, ref)

    def test_secao_1_identificacao(self):
        ger = dict(_linhas(self.ger["1"][0]))
        ref = dict(_linhas(self.ref["1"][0]))
        ref["Parceiro(s)"] = ref.pop("Parceiro principal")
        for r in ("Data de geração",):
            ger.pop(r), ref.pop(r)
        self.assertEqual(ger, ref)
        self.assertEqual(dict(_linhas(self.ger["1"][0]))["Data de geração"],
                         "15/10/2026 — 21:00")

    def test_secao_3_visao_geral(self):
        linhas = _linhas(self.ger["3"][0])
        self.assertEqual(linhas[0], ("Elemento analisado", "Informação registrada"))
        self.assertEqual([l[0] for l in linhas[1:]], [
            "ID do registro", "Ator focal", "Tipo de ator", "Parceiro(s)",
            "Atividade executada/documentada", "Mecanismo de aprendizagem observado",
            "Evidência de aplicação tecnológica",
            "Relação com trajetória de CT da função", "Status da análise"])

    def test_secao_4_vinculo(self):
        self.comparar("4")
        self.assertEqual(_caixa(self.ger["4"][1]), _norm(
            textos.LEITURA_GERENCIAL_FIXA + " " + textos.LEITURA_GERENCIAL_POR_GRAU["A"]))

    def test_secao_5_mecanismo(self):
        self.comparar("5")

    def test_secao_6_fluxo(self):
        self.comparar("6", ignorar_rotulos=("Variedade observável",))
        self.assertEqual(_caixa(self.ger["6"][1]), _caixa(self.ref["6"][1]))

    def test_secao_7_aplicacao(self):
        self.comparar("7")

    def test_secao_8_trajetoria(self):
        self.comparar("8", so_valores=True)
        self.assertEqual(_caixa(self.ger["8"][1]), _caixa(self.ref["8"][1]))

    def test_secao_9_quadro(self):
        self.comparar("9")

    def test_secao_12_informacao_adicional(self):
        self.comparar("12", ignorar_rotulos=("Condição de preenchimento",))

    def test_nenhum_marcador_sobra(self):
        d = docx.Document(str(self.saida))
        textos_doc = [p.text for p in d.paragraphs]
        textos_doc += [c.text for t in d.tables for r in t.rows for c in r.cells]
        textos_doc += [p.text for p in d.sections[0].footer.paragraphs]
        self.assertFalse([t for t in textos_doc if "{{" in t or "}}" in t])


if __name__ == "__main__":
    unittest.main()
