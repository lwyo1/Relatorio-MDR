import tempfile
import unittest
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook

from mdr import leitura, montagem
from mdr import planilha as P
from mdr.campos import CAMPOS
from mdr.estados import como_texto
from tests import dados_gripen as G

PLANILHA = Path(__file__).resolve().parent.parent / "modelos" / "planilha_mdr.xlsx"
AGORA = datetime(2026, 10, 15, 21, 0)


def _t(v) -> str:
    return como_texto(v).strip()


class TestPlanilhaModelo(unittest.TestCase):
    """A planilha de entrada gerada deve conter exatamente o caso Gripen."""

    @classmethod
    def setUpClass(cls):
        cls.base = leitura.ler(PLANILHA)
        cls.esperada = G.base_gripen()

    def test_sem_avisos(self):
        self.assertEqual(self.base.avisos, [])

    def test_registro_igual_ao_gripen(self):
        lido = self.base.registro("MDR-001").valores
        for k in CAMPOS:
            self.assertEqual(_t(lido.get(k)), _t(G.VALORES_MDR001.get(k)), k)

    def test_tabelas_auxiliares(self):
        self.assertEqual(self.base.fontes, self.esperada.fontes)
        self.assertEqual(self.base.limitacoes, self.esperada.limitacoes)
        self.assertEqual(self.base.lacunas, self.esperada.lacunas)
        self.assertEqual(self.base.fases, self.esperada.fases)
        self.assertEqual(self.base.trajetorias, self.esperada.trajetorias)
        self.assertEqual([a.id_acao for a in self.base.acoes], ["ACAO-001"])
        self.assertEqual(self.base.acoes[0].lacunas, ["LAC-001", "LAC-002", "LAC-003"])
        for k, v in self.esperada.acoes[0].campos.items():
            self.assertEqual(_t(self.base.acoes[0].campos[k]), v, k)

    def test_relatorio_da_planilha_igual_ao_dos_dados(self):
        c1 = montagem.montar(self.base.registro("MDR-001"), self.base, AGORA)
        c2 = montagem.montar(self.esperada.registro("MDR-001"), self.esperada, AGORA)
        self.assertEqual(c1, c2)


class TestPlanilhaComProblemas(unittest.TestCase):
    """Planilhas montadas à mão com defeitos de estrutura."""

    def _salvar(self, wb) -> Path:
        self.tmp = tempfile.TemporaryDirectory()
        caminho = Path(self.tmp.name) / "p.xlsx"
        wb.save(str(caminho))
        return caminho

    def tearDown(self):
        if hasattr(self, "tmp"):
            self.tmp.cleanup()

    def test_colunas_em_outra_ordem_e_coluna_ausente(self):
        wb = load_workbook(str(PLANILHA))
        ws = wb[P.REGISTROS]
        ws.delete_cols(ws.max_column)          # tira "Síntese das limitações..."
        ws.move_range("A2:A3", rows=0, cols=ws.max_column)  # ID vai para o fim
        base = leitura.ler(self._salvar(wb))
        self.assertEqual(base.registro("MDR-001").get("programa"), G.PROGRAMA)
        self.assertTrue(any("Síntese das limitações" in a for a in base.avisos))

    def test_id_repetido_e_linha_sem_id(self):
        wb = load_workbook(str(PLANILHA))
        ws = wb[P.REGISTROS]
        for col in range(1, ws.max_column + 1):
            ws.cell(4, col, ws.cell(3, col).value)   # duplica MDR-001
        ws.cell(5, 2, "Programa sem ID")
        base = leitura.ler(self._salvar(wb))
        self.assertEqual(len(base.registros), 1)
        self.assertTrue(any("repetido" in a for a in base.avisos))
        self.assertTrue(any("falta o ID" in a for a in base.avisos))

    def test_evidencia_e_lacuna_com_id_inexistente(self):
        wb = load_workbook(str(PLANILHA))
        wb[P.FONTES].append(["EVD-009", "MDR-999", "Fonte X", "Evidência X", "Entrevista"])
        wb[P.LACUNAS].append(["LAC-009", "Lacuna X", "MDR-001; MDR-998"])
        base = leitura.ler(self._salvar(wb))
        self.assertTrue(any("MDR-999" in a for a in base.avisos))
        self.assertTrue(any("MDR-998" in a for a in base.avisos))
        self.assertEqual(base.lacunas[-1].registros, ["MDR-001", "MDR-998"])

    def test_sem_aba_registros(self):
        wb = Workbook()
        wb.active.title = "Outra"
        with self.assertRaises(leitura.ErroPlanilha):
            leitura.ler(self._salvar(wb))

    def test_arquivo_inexistente(self):
        with self.assertRaises(leitura.ErroPlanilha):
            leitura.ler("/caminho/que/nao/existe.xlsx")


if __name__ == "__main__":
    unittest.main()
