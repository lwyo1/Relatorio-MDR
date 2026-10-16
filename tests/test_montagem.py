import tempfile
import unittest
from pathlib import Path

from mdr import gerador, montagem
from mdr.estados import NAO_PREENCHIDO
from mdr.modelos import Base, Registro
from tests.dados_gripen import VALORES_MDR001, base_gripen

TEMPLATE = Path(__file__).resolve().parent.parent / "modelos" / "relatorio_mdr.docx"


def _registro(**alteracoes) -> tuple[Registro, Base]:
    base = base_gripen()
    valores = dict(VALORES_MDR001, **alteracoes)
    base.registros = [Registro(valores)]
    return base.registros[0], base


class TestProibicoes(unittest.TestCase):
    def test_atividade_prevista_nunca_preenche_a_executada(self):
        r, b = _registro(atividade_executada="")
        c = montagem.montar(r, b)
        self.assertEqual(c["atividade_executada"], NAO_PREENCHIDO)
        self.assertNotEqual(c["atividade_executada"], c["atividade_prevista"])

    def test_funcao_associada_nao_copia_funcao_do_recorte(self):
        r, b = _registro(funcao_associada=None)
        self.assertEqual(montagem.montar(r, b)["funcao_associada"], NAO_PREENCHIDO)

    def test_origem_nao_vem_da_direcao(self):
        r, b = _registro(origem_conhecimento="Não determinado")
        self.assertEqual(montagem.montar(r, b)["origem_conhecimento"],
                         "Não determinada com as evidências disponíveis")

    def test_periodo_e_trajetoria_so_por_cadastro(self):
        r, b = _registro(fase_temporal="Fase 9 — Inexistente")
        c = montagem.montar(r, b)
        self.assertEqual(c["periodo"], NAO_PREENCHIDO)
        self.assertEqual(c["trajetoria"], NAO_PREENCHIDO)

    def test_relacao_ct_nao_e_calculada_dos_graus(self):
        r, b = _registro(relacao_ct="Relação indeterminada", grau_aplicacao="A")
        self.assertEqual(montagem.montar(r, b)["relacao_ct"], "Relação indeterminada")


class TestGerador(unittest.TestCase):
    def test_marcador_sem_conteudo_e_erro(self):
        r, b = _registro()
        c = montagem.montar(r, b)
        del c["mecanismo"]
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaises(gerador.MarcadorSemConteudo):
                gerador.gerar(TEMPLATE, c, Path(tmp) / "x.docx")

    def test_linhas_dinamicas_seguem_a_quantidade_de_itens(self):
        import docx
        r, b = _registro()
        c = montagem.montar(r, b)
        c["fontes"] = [(f"Fonte {i}", f"Evidência {i}", "Documental") for i in range(8)]
        with tempfile.TemporaryDirectory() as tmp:
            saida = gerador.gerar(TEMPLATE, c, Path(tmp) / "x.docx")
            d = docx.Document(str(saida))
            tabela = next(t for t in d.tables
                          if t.cell(0, 0).text.startswith("Fonte/Referência"))
            self.assertEqual(len(tabela.rows), 1 + 8)
            self.assertEqual(tabela.cell(8, 0).text, "Fonte 7")


if __name__ == "__main__":
    unittest.main()
