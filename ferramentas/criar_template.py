"""Cria modelos/relatorio_mdr.docx a partir do relatório-modelo do Gripen.

O visual (cores, fontes, tabelas, rodapé) vem do modelo preenchido do Gripen.
Este script troca cada conteúdo específico do Gripen por um marcador {{chave}}
e aplica os ajustes da especificação (decisões D5, D11, D12 e o rótulo genérico
do bloco 8). Rodar de novo sempre que o template precisar mudar:

    python ferramentas/criar_template.py
"""

import copy
import sys
from pathlib import Path

import docx
from docx.oxml.ns import qn

RAIZ = Path(__file__).resolve().parent.parent
ORIGEM = RAIZ / "modelos" / "modelo_gripen_referencia.docx"
DESTINO = RAIZ / "modelos" / "relatorio_mdr.docx"


def _limpar_celula(cell, manter_runs=1):
    """Mantém só o primeiro parágrafo e os primeiros runs (com sua formatação)."""
    for p in cell.paragraphs[1:]:
        p._p.getparent().remove(p._p)
    p = cell.paragraphs[0]
    for r in p.runs[manter_runs:]:
        r._r.getparent().remove(r._r)
    return p


def valor(cell, texto):
    p = _limpar_celula(cell, 1)
    p.runs[0].text = texto


def rotulo(cell, texto):
    valor(cell, texto)


def caixa(tabela, conteudo, titulo=None):
    """Caixa destacada (tabela 1x1): run 0 = título, run 1 = conteúdo."""
    p = _limpar_celula(tabela.cell(0, 0), 2)
    if titulo is not None:
        p.runs[0].text = titulo + "\n"
    p.runs[1].text = conteudo


def coluna_valores(tabela, chaves, col=1):
    """Tabela chave-valor: troca a coluna de valores pelos marcadores, linha a linha."""
    assert len(tabela.rows) == len(chaves), (len(tabela.rows), chaves)
    for row, chave in zip(tabela.rows, chaves):
        valor(row.cells[col], "{{" + chave + "}}")


def tabela_dinamica(tabela, nome, ncols, cabecalho=None, manter_finais=0):
    """Cabeçalho + 1 linha-modelo {{nome.i}} (repetida pelo gerador) + linhas finais fixas."""
    linhas = list(tabela.rows)
    if cabecalho:
        for cell, texto in zip(linhas[0].cells, cabecalho):
            rotulo(cell, texto)
    modelo = linhas[1]
    for i in range(ncols):
        valor(modelo.cells[i], "{{" + f"{nome}.{i}" + "}}")
    fim = len(linhas) - manter_finais
    for row in linhas[2:fim]:
        row._tr.getparent().remove(row._tr)


def paragrafo_rotulado(p, chave):
    """'Rótulo: ' em negrito + conteúdo -> mantém o rótulo e troca o conteúdo."""
    for r in p.runs[2:]:
        r._r.getparent().remove(r._r)
    p.runs[1].text = "{{" + chave + "}}"


def paragrafo_simples(p, chave):
    for r in p.runs[1:]:
        r._r.getparent().remove(r._r)
    p.runs[0].text = "{{" + chave + "}}"


# Ordem dos filhos de <w:trPr> exigida pelo esquema do Word.
_ORDEM_TRPR = ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
               "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden",
               "ins", "del", "trPrChange"]


def _propriedade_linha(row, nome):
    trPr = row._tr.get_or_add_trPr()
    if trPr.find(qn(f"w:{nome}")) is not None:
        return
    novo = trPr.makeelement(qn(f"w:{nome}"), {})
    pos = _ORDEM_TRPR.index(nome)
    for filho in trPr:
        local = filho.tag.split("}")[1]
        if local in _ORDEM_TRPR and _ORDEM_TRPR.index(local) > pos:
            filho.addprevious(novo)
            return
    trPr.append(novo)


def criar():
    d = docx.Document(str(ORIGEM))
    t = d.tables
    p = d.paragraphs

    # Título
    paragrafo_simples(p[1], "id_registro")
    p[1].runs[0].text = "RELATÓRIO DE ANÁLISE DO REGISTRO {{id_registro}}"
    p[2].runs[0].text = "{{programa}}\n{{funcao_tecnologica}} | {{fase_temporal}}"
    for r in p[2].runs[1:]:
        r._r.getparent().remove(r._r)
    caixa(t[0], "{{texto_natureza}}")

    # 1. Identificação (D12: "Parceiro principal" -> "Parceiro(s)")
    rotulo(t[1].rows[8].cells[0], "Parceiro(s)")
    coluna_valores(t[1], [
        "id_registro", "programa", "fase_temporal", "funcao_tecnologica", "periodo",
        "quantidade_registros", "ator_focal", "tipo_ator", "parceiros",
        "status_analise", "data_geracao", "responsavel",
    ])

    # 2. Resumo executivo
    resumo = next(x for x in p if x.text.startswith("O registro MDR-001 examina"))
    paragrafo_simples(resumo, "resumo_executivo")

    # 3. Visão geral (D5: tabela vertical de 2 colunas, como na regra de geração).
    #    Reaproveita o visual da tabela de limitações (cabeçalho escuro + listras).
    nova = copy.deepcopy(t[13]._tbl)
    t[2]._tbl.addprevious(nova)
    t[2]._tbl.getparent().remove(t[2]._tbl)
    t = d.tables  # índices atualizados (a nova tabela ocupa a posição 2)
    tabela_dinamica(t[2], "visao_geral", 2,
                    cabecalho=["Elemento analisado", "Informação registrada"])

    # 4. Vínculo + Leitura gerencial
    coluna_valores(t[3], ["natureza_vinculo", "atividade_prevista",
                          "atividade_executada", "grau_atividade_extenso"])
    caixa(t[4], "{{leitura_gerencial}}")

    # 5. Mecanismo
    coluna_valores(t[5], ["mecanismo", "conhecimento", "tipo_mecanismo",
                          "grau_mecanismo_extenso"])

    # 6. Fluxo (D11: sem "Variedade observável" no relatório individual)
    variedade = t[6].rows[3]
    assert variedade.cells[0].text.startswith("Variedade"), variedade.cells[0].text
    variedade._tr.getparent().remove(variedade._tr)
    coluna_valores(t[6], ["parceiro_fluxo", "direcao_fluxo", "origem_conhecimento",
                          "impactos", "grau_fluxo_extenso"])
    caixa(t[7], "{{limitacao_fluxo}}")

    # 7. Aplicação tecnológica
    coluna_valores(t[8], ["evidencia_aplicacao", "atividade_tecnologica",
                          "funcao_associada", "descritor", "grau_aplicacao_extenso"])

    # 8. Trajetória (rótulo genérico, item 8.1)
    rotulo(t[9].rows[0].cells[0], "Trajetória da função no período/fase analisada")
    coluna_valores(t[9], ["trajetoria", "relacao_ct", "grau_ct_extenso",
                          "interpretacao_ct"])
    caixa(t[10], "{{limitacao_ct}}")

    # 9. Quadro de força + nota fixa
    tabela_dinamica(t[11], "quadro", 3)
    nota = next(x for x in p if x.text.startswith("Nota metodológica"))
    paragrafo_rotulado(nota, "nota_metodologica")

    # 10. Fontes + observação fixa
    tabela_dinamica(t[12], "fontes", 3, cabecalho=[
        "Fonte/Referência no projeto", "Evidência utilizada no registro", "Tipo"])
    obs = next(x for x in p if x.text.startswith("Observação"))
    paragrafo_rotulado(obs, "observacao_fontes")

    # 11. Limitações (linhas dinâmicas) + linha final de alertas
    tabela_dinamica(t[13], "limitacoes", 2, manter_finais=1)
    valor(t[13].rows[-1].cells[1], "{{alertas}}")

    # 12. Informação adicional
    coluna_valores(t[14], ["informacao_adicional", "fonte_complementar",
                           "condicao_preenchimento"])

    # 13. Diagnóstico
    coluna_valores(t[15], ["diagnostico_sugerido", "fundamento_sugestao",
                           "diagnostico_confirmado_txt", "houve_alteracao",
                           "justificativa_revisao_txt"])
    caixa(t[16], "{{implicacao_gerencial}}", titulo="{{titulo_implicacao}}")

    # 14. Lacunas
    tabela_dinamica(t[17], "lacunas", 3)

    # 15. Plano de ação (uma ação por enquanto; várias ações na etapa E3)
    valor(t[18].rows[0].cells[1], "{{titulo_acao}}")
    chaves_acao = ["problema", "objetivo", "acao", "justificativa",
                   "responsavel_potencial", "atores", "momento", "procedimento",
                   "indicador", "evidencia_esperada", "prazo", "risco"]
    for row, k in zip(t[18].rows[1:], chaves_acao):
        valor(row.cells[1], "{{acao_" + k + "}}")

    # 16. Conclusão
    conclusao = next(x for x in p if x.text.startswith("O registro MDR-001 permite"))
    paragrafo_simples(conclusao, "conclusao")

    # Paginação: nenhuma linha de tabela (nem caixa) se parte entre páginas, e o
    # cabeçalho das tabelas com cabeçalho se repete em nova página (item 3).
    com_cabecalho = [t[2], t[11], t[12], t[13], t[17], t[18]]
    for tabela in d.tables:
        for i, row in enumerate(tabela.rows):
            _propriedade_linha(row, "cantSplit")
            if i == 0 and any(tabela._tbl is x._tbl for x in com_cabecalho):
                _propriedade_linha(row, "tblHeader")

    # Rodapé (o campo PAGE do número da página é preservado)
    rod = d.sections[0].footer.paragraphs[0]
    rod.runs[0].text = ("{{id_registro}} | {{programa}} | {{funcao_tecnologica}} | "
                        "{{fase_temporal}}   •   Página ")

    d.save(str(DESTINO))
    return DESTINO


if __name__ == "__main__":
    print("Template criado em", criar())
    sys.exit(0)
