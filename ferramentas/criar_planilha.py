"""Cria modelos/planilha_mdr.xlsx: a planilha de entrada da MDR-BID Digital.

A estrutura vem de mdr/planilha.py (a mesma que a leitura usa). A linha de
exemplo é o caso MDR-001 (Gripen). Rodar de novo sempre que a estrutura mudar:

    python ferramentas/criar_planilha.py
"""

import sys
from pathlib import Path

RAIZ = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(RAIZ))

from openpyxl import Workbook  # noqa: E402
from openpyxl.comments import Comment  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.utils import get_column_letter  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

from mdr import planilha as P  # noqa: E402
from mdr.campos import CAMPOS  # noqa: E402
from tests import dados_gripen as G  # noqa: E402

DESTINO = RAIZ / "modelos" / "planilha_mdr.xlsx"
LINHAS_COM_LISTA = 1000

AZUL = "17365D"
FONTE = Font(name="Arial", size=10)
FONTE_CAB = Font(name="Arial", size=10, bold=True, color="FFFFFF")
FONTE_GRUPO = Font(name="Arial", size=10, bold=True, color=AZUL)
FUNDO_CAB = PatternFill("solid", fgColor=AZUL)
FUNDO_GRUPO = PatternFill("solid", fgColor="D9EAF7")
FUNDO_EXEMPLO = PatternFill("solid", fgColor="FFF2CC")
BORDA = Border(*(Side(style="thin", color="BFBFBF"),) * 4)
QUEBRA = Alignment(wrap_text=True, vertical="top")

LARGURA_CURTA = {"id_registro", "grau_atividade", "grau_mecanismo", "grau_fluxo",
                 "grau_aplicacao", "grau_ct", "diagnostico_confirmado"}

DICAS = {
    "id_registro": "Padrão MDR-001, MDR-002...",
    "grau_atividade": "Só a letra: A, B, C ou D (ou Não aplicável).",
    "fase_temporal": "Use exatamente o nome cadastrado na aba FASES.",
    "status_analise": "Decisão da pesquisadora: o sistema nunca marca 'Concluída' sozinho.",
    "responsavel": "Nome de quem analisou e validou o registro.",
    "limitacao_fluxo": "Texto da caixa 'Limite da caracterização do fluxo' (bloco 6).",
    "interpretacao_ct": "Interpretação validada pela pesquisadora (bloco 8).",
    "limitacao_ct": "Texto da caixa 'Limite interpretativo' (bloco 8).",
    "diagnostico_confirmado": "Diagnóstico escolhido pelo responsável (I a IV).",
    "sintese_registro": "Uma frase com o principal encadeamento, para a conclusão.",
    "sintese_limitacoes": "Síntese curta das limitações, para o resumo e a conclusão.",
}
for g in ("grau_mecanismo", "grau_fluxo", "grau_aplicacao", "grau_ct"):
    DICAS[g] = DICAS["grau_atividade"]


def _cabecalho(ws, linha, textos):
    for col, texto in enumerate(textos, 1):
        c = ws.cell(linha, col, texto)
        c.font, c.fill, c.alignment, c.border = FONTE_CAB, FUNDO_CAB, QUEBRA, BORDA


def _linha(ws, linha, valores, exemplo=False):
    for col, v in enumerate(valores, 1):
        c = ws.cell(linha, col, v if v != "" else None)
        c.font, c.alignment, c.border = FONTE, QUEBRA, BORDA
        if exemplo:
            c.fill = FUNDO_EXEMPLO


def _larguras(ws, larguras):
    for col, w in enumerate(larguras, 1):
        ws.column_dimensions[get_column_letter(col)].width = w


def _lista(ws, coluna_letra, nome_lista, linha_inicial, listas_ref):
    dv = DataValidation(type="list", formula1=listas_ref[nome_lista], allow_blank=True,
                        errorStyle="warning", showErrorMessage=True,
                        errorTitle="Valor fora da lista",
                        error="Este valor não está na lista. Mantenha só se for um "
                              "texto próprio do registro.")
    dv.add(f"{coluna_letra}{linha_inicial}:{coluna_letra}{LINHAS_COM_LISTA}")
    ws.add_data_validation(dv)


def aba_listas(wb):
    ws = wb.create_sheet(P.LISTAS)
    refs = {}
    for col, (nome, opcoes) in enumerate(P.OPCOES.items(), 1):
        letra = get_column_letter(col)
        c = ws.cell(1, col, nome)
        c.font, c.fill, c.alignment = FONTE_CAB, FUNDO_CAB, QUEBRA
        for i, op in enumerate(opcoes, 2):
            ws.cell(i, col, op).font = FONTE
        refs[nome] = f"={P.LISTAS}!${letra}$2:${letra}${len(opcoes) + 1}"
        ws.column_dimensions[letra].width = 32
    ws.freeze_panes = "A2"
    return refs


def aba_registros(wb, refs):
    ws = wb.create_sheet(P.REGISTROS)
    col = 1
    for grupo, chaves in P.GRUPOS_REGISTROS:
        ws.merge_cells(start_row=1, start_column=col, end_row=1,
                       end_column=col + len(chaves) - 1)
        c = ws.cell(1, col, grupo)
        c.font, c.fill, c.alignment = FONTE_GRUPO, FUNDO_GRUPO, QUEBRA
        col += len(chaves)
    _cabecalho(ws, 2, [CAMPOS[k].rotulo for k in P.COLUNAS_REGISTROS])
    for i, k in enumerate(P.COLUNAS_REGISTROS, 1):
        if k in DICAS:
            ws.cell(2, i).comment = Comment(DICAS[k], "MDR-BID")
    _linha(ws, 3, [G.VALORES_MDR001.get(k, "") for k in P.COLUNAS_REGISTROS],
           exemplo=True)
    _larguras(ws, [14 if k in LARGURA_CURTA else 34 for k in P.COLUNAS_REGISTROS])
    for i, k in enumerate(P.COLUNAS_REGISTROS, 1):
        letra = get_column_letter(i)
        if k in P.LISTA_DA_COLUNA:
            _lista(ws, letra, P.LISTA_DA_COLUNA[k], 3, refs)
        if k == "fase_temporal":
            dv = DataValidation(type="list", formula1=f"={P.FASES}!$B$2:$B$200",
                                allow_blank=True, errorStyle="warning",
                                showErrorMessage=True, errorTitle="Fase não cadastrada",
                                error="Cadastre a fase na aba FASES para o período "
                                      "aparecer no relatório.")
            dv.add(f"{letra}3:{letra}{LINHAS_COM_LISTA}")
            ws.add_data_validation(dv)
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 42
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(P.COLUNAS_REGISTROS))}2"


def aba_auxiliar(wb, nome, linhas, larguras, listas=None, refs=None, dicas=None):
    ws = wb.create_sheet(nome)
    colunas = P.AUXILIARES[nome]
    _cabecalho(ws, 1, [cab for _, cab in colunas])
    for i, (atributo, _) in enumerate(colunas, 1):
        if dicas and atributo in dicas:
            ws.cell(1, i).comment = Comment(dicas[atributo], "MDR-BID")
    for n, valores in enumerate(linhas, 2):
        _linha(ws, n, valores, exemplo=True)
    _larguras(ws, larguras)
    for atributo, nome_lista in (listas or {}).items():
        idx = [a for a, _ in colunas].index(atributo) + 1
        _lista(ws, get_column_letter(idx), nome_lista, 2, refs)
    ws.freeze_panes = "A2"
    return ws


def aba_leia_me(wb):
    ws = wb.active
    ws.title = "LEIA-ME"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    linhas = [
        ("MDR-BID Digital — Planilha de entrada", None, "titulo"),
        ("Para que serve", "Cada linha da aba REGISTROS é um registro MDR. O programa "
         "lê esta planilha e gera o relatório em Word do registro escolhido. O "
         "relatório só reproduz o que está aqui: nada é completado ou interpretado.",
         None),
        ("Linha de exemplo", "As linhas com fundo amarelo são o caso MDR-001 (Programa "
         "F-X2 Gripen), copiado do relatório-modelo. Servem de exemplo de formato e "
         "podem ser apagadas.", None),
        ("Como preencher", None, "secao"),
        ("REGISTROS", "Uma linha por registro. As colunas seguem os blocos do "
         "relatório (linha 1). Passe o mouse no cabeçalho para ver dicas.", None),
        ("FONTES", "Uma linha por evidência, ligada ao registro pelo ID MDR. Um "
         "registro pode ter várias; a mesma fonte pode aparecer em várias linhas.",
         None),
        ("LIMITACOES", "Uma linha por limitação/indeterminação do registro (bloco 11).",
         None),
        ("LACUNAS", "Uma linha por lacuna (LAC-001...). Em 'Registros relacionados', "
         "os IDs MDR separados por ponto e vírgula: MDR-001; MDR-003.", None),
        ("PLANO_ACAO", "Uma linha por ação (ACAO-001...), com os registros e lacunas "
         "relacionados separados por ponto e vírgula.", None),
        ("FASES", "Período de cada fase de cada programa (bloco 1).", None),
        ("TRAJETORIA", "Trajetória/nível validado de cada função, por programa e fase "
         "(bloco 8). O sistema não calcula o nível.", None),
        ("LISTAS", "Opções das listas suspensas. PROVISÓRIAS até a confirmação das "
         "listas oficiais.", None),
        ("Regras de preenchimento", None, "secao"),
        ("Não determinado", "Escreva exatamente 'Não determinado' quando você analisou "
         "e as evidências não bastam. Aparece no relatório como 'Não determinado com "
         "as evidências disponíveis'.", None),
        ("Não aplicável", "Escreva exatamente 'Não aplicável' quando a dimensão não se "
         "aplica. Aparece como 'Não aplicável ao registro analisado'.", None),
        ("Célula vazia", "Significa 'ainda não preenchido'. Gera alerta e nunca é "
         "convertida em 'Não determinado' nem em grau D.", None),
        ("Graus", "Só a letra: A, B, C ou D. O relatório acrescenta a legenda "
         "(A — Evidência sustentada etc.). Não há média nem grau geral.", None),
        ("Listas suspensas", "Um valor fora da lista mostra um aviso, mas pode ser "
         "mantido quando for um texto próprio do registro.", None),
    ]
    for n, (a, b, tipo) in enumerate(linhas, 1):
        ca, cb = ws.cell(n, 1, a), ws.cell(n, 2, b)
        if tipo == "titulo":
            ca.font = Font(name="Arial", size=16, bold=True, color=AZUL)
            ws.row_dimensions[n].height = 26
        elif tipo == "secao":
            ca.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            ca.fill = cb.fill = FUNDO_CAB
        else:
            ca.font = Font(name="Arial", size=10, bold=True, color=AZUL)
            cb.font = FONTE
        ca.alignment = cb.alignment = QUEBRA
        if tipo == "titulo":
            ca.alignment = Alignment(vertical="top")  # título sem quebra de linha
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


def criar():
    wb = Workbook()
    aba_leia_me(wb)
    # LISTAS é criada antes para obter as referências, e movida para o fim depois.
    refs = aba_listas(wb)
    aba_registros(wb, refs)
    aba_auxiliar(wb, P.FONTES,
                 [[f.id_evidencia, f.id_registro, f.fonte, f.evidencia, f.tipo]
                  for f in G.FONTES], [16, 12, 40, 60, 30],
                 listas={"tipo": "Tipo de evidência"}, refs=refs)
    aba_auxiliar(wb, P.LIMITACOES,
                 [[x.id_limitacao, x.id_registro, x.elemento, x.situacao]
                  for x in G.LIMITACOES], [16, 12, 36, 80])
    aba_auxiliar(wb, P.LACUNAS,
                 [[x.id_lacuna, x.descricao, "; ".join(x.registros)]
                  for x in G.LACUNAS], [14, 80, 28],
                 dicas={"registros": "IDs MDR separados por ponto e vírgula."})
    aba_auxiliar(wb, P.PLANO_ACAO,
                 [[a.id_acao, "; ".join(a.registros), "; ".join(a.lacunas)]
                  + [a.campos.get(k, "") for k, _ in P.CAMPOS_ACAO]
                  for a in G.ACOES], [14, 18, 24] + [40] * len(P.CAMPOS_ACAO),
                 dicas={"registros": "IDs MDR separados por ponto e vírgula.",
                        "lacunas": "IDs de lacuna separados por ponto e vírgula."})
    aba_auxiliar(wb, P.FASES,
                 [[prog, fase, per] for (prog, fase), per in G.base_gripen().fases.items()],
                 [32, 48, 16])
    aba_auxiliar(wb, P.TRAJETORIA,
                 [[prog, fase, func, traj] for (prog, fase, func), traj
                  in G.base_gripen().trajetorias.items()], [28, 44, 28, 60])
    wb.move_sheet(P.LISTAS, offset=len(wb.sheetnames))
    wb.active = 0
    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(DESTINO))
    return DESTINO


if __name__ == "__main__":
    print("Planilha criada em", criar())
